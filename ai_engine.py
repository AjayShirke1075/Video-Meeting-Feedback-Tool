"""
InterviewIQ - AI Engine
Handles all Gemini AI operations: text extraction, question generation, answer analysis
"""
from __future__ import annotations

import os
import json
import re
import time
import logging
import config
from emotion_engine import fuse_biometrics, adjust_answer_score

config.load_env()
logger = logging.getLogger(__name__)

# ─────────────────────────────────────────────
#  Gemini Setup  (google-genai SDK)
# ─────────────────────────────────────────────
try:
    from google import genai
    from google.genai import types as genai_types
    _GENAI_IMPORTED = True
except ImportError:
    _GENAI_IMPORTED = False

import socket

# Force IPv4 socket resolution to prevent blackholed IPv6 dual-stack timeouts on Windows
try:
    _orig_getaddrinfo = socket.getaddrinfo
    def _ipv4_getaddrinfo(host, port, family=0, type=0, proto=0, flags=0):
        return _orig_getaddrinfo(host, port, socket.AF_INET, type, proto, flags)
    socket.getaddrinfo = _ipv4_getaddrinfo
except Exception:
    pass

# Primary model (override with env var GEMINI_MODEL). Fallbacks are tried on 404/503.
_MODEL = os.getenv("GEMINI_MODEL", "gemini-3.5-flash-lite")
_FALLBACK_MODELS = [
    "gemini-3.5-flash-lite",
    "gemini-3.1-flash-lite",
    "gemini-flash-lite-latest",
    "gemini-2.5-flash",
    "gemini-2.5-flash-lite",
]


def _model_candidates() -> list[str]:
    seen, out = set(), []
    for m in [_MODEL] + _FALLBACK_MODELS:
        if m and m not in seen:
            seen.add(m)
            out.append(m)
    return out


_GEMINI_READY = False
_gemini_client = None
_client_key = None


def _load_api_key() -> str:
    return config.get_api_key()


def _is_bad_key_error(msg: str) -> bool:
    m = msg.lower()
    return ("api key not valid" in m or "api_key_invalid" in m
            or "permission_denied" in m or "unauthenticated" in m
            or " 401" in m or "403" in m)


def _is_model_not_found(msg: str) -> bool:
    m = msg.lower()
    return "404" in m or "not_found" in m or "no longer available" in m


def _is_transient_or_model_error(msg: str) -> bool:
    m = msg.lower()
    return (_is_model_not_found(m)
            or "503" in m or "unavailable" in m or "high demand" in m
            or "504" in m or "deadline_exceeded" in m or "timeout" in m or "timed out" in m
            or "resource_exhausted" in m or "429" in m or "quota" in m or "rate limit" in m)


def _is_permanent_error(msg: str) -> bool:
    return _is_bad_key_error(msg) or "invalid_argument" in msg.lower()


def save_api_key(api_key: str) -> bool:
    """Save and activate a new Gemini API key."""
    config.save_api_key(api_key)
    return _configure_gemini(force=True)


def validate_api_key(api_key: str) -> tuple[bool, str]:
    """Test an API key. Returns (success, message)."""
    if not _GENAI_IMPORTED:
        return False, "google-genai package not installed. Run: pip install google-genai"
    key = api_key.strip()
    if not key or key in config.PLACEHOLDER_KEYS:
        return False, "API key is empty or a placeholder - enter a valid key"

    client = genai.Client(
        api_key=key,
        http_options=genai_types.HttpOptions(timeout=12000),
    )

    # 1. Verify key authenticity with Google API
    try:
        next(iter(client.models.list()), None)
    except Exception as exc:
        err = str(exc)
        if _is_bad_key_error(err):
            return False, "API key is invalid. Please check it and try again."
        return False, f"Could not verify key: {err}"

    # 2. Test generation across candidate models
    last_error = ""
    for model in _model_candidates():
        try:
            resp = client.models.generate_content(
                model=model,
                contents="Reply with exactly: OK",
            )
            if resp.text:
                return True, "API key verified"
            last_error = "Empty response from Gemini"
        except Exception as exc:
            last_error = str(exc)
            if _is_bad_key_error(last_error):
                return False, "API key is invalid. Please check it and try again."
            if _is_transient_or_model_error(last_error):
                logger.info("Model %s unavailable during check, trying next: %s", model, last_error)
                continue
            return False, f"Could not verify key: {last_error}"

    # Key is valid and authenticated, models are temporarily in high demand
    return True, "API key verified"


def _configure_gemini(force: bool = False) -> bool:
    """Build (or reuse) a Gemini client from the current key. Returns True if ready."""
    global _GEMINI_READY, _gemini_client, _client_key
    if not _GENAI_IMPORTED:
        _GEMINI_READY = False
        return False
    api_key = _load_api_key()
    if not api_key or api_key in config.PLACEHOLDER_KEYS:
        _GEMINI_READY = False
        _gemini_client = None
        _client_key = None
        return False
    if not force and _gemini_client is not None and _client_key == api_key:
        return True
    try:
        _gemini_client = genai.Client(
            api_key=api_key,
            http_options=genai_types.HttpOptions(timeout=30000),
        )
        _client_key = api_key
        _GEMINI_READY = True
        return True
    except Exception as exc:
        logger.warning("Failed to create Gemini client: %s", exc)
        _GEMINI_READY = False
        return False


_configure_gemini()

# ─────────────────────────────────────────────
#  PDF / DOCX / TXT Text Extraction
# ─────────────────────────────────────────────
def extract_text(filepath: str, filename: str) -> str:
    """Extract plain text from PDF, DOCX, or TXT files."""
    ext = filename.rsplit(".", 1)[-1].lower() if "." in filename else "txt"
    if ext == "pdf":
        text = _extract_pdf(filepath)
    elif ext == "docx":
        text = _extract_docx(filepath)
    else:
        with open(filepath, "r", encoding="utf-8", errors="ignore") as f:
            text = f.read()
    if not text.strip():
        raise RuntimeError(
            "No text could be extracted from the file (it may be a scanned image)."
        )
    return text


def _extract_pdf(filepath: str) -> str:
    try:
        import fitz  # PyMuPDF
        doc = fitz.open(filepath)
        text = "\n".join(page.get_text() for page in doc)
        doc.close()
        return text.strip()
    except ImportError:
        logger.warning("PyMuPDF not installed, trying fallback PDF parser")
        return _extract_pdf_fallback(filepath)
    except Exception as e:
        raise RuntimeError(f"Failed to parse PDF: {e}")


def _extract_pdf_fallback(filepath: str) -> str:
    """Fallback PDF parser using PyPDF2."""
    try:
        import PyPDF2
        with open(filepath, "rb") as f:
            reader = PyPDF2.PdfReader(f)
            return "\n".join(page.extract_text() or "" for page in reader.pages).strip()
    except Exception as e:
        raise RuntimeError(f"PDF parsing failed: {e}")


def _extract_docx(filepath: str) -> str:
    try:
        from docx import Document
        doc = Document(filepath)
        parts = [p.text for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                cells = [c.text.strip() for c in row.cells if c.text.strip()]
                if cells:
                    parts.append(" | ".join(cells))
        return "\n".join(parts).strip()
    except ImportError:
        raise RuntimeError("python-docx not installed. Run: pip install python-docx")
    except Exception as e:
        raise RuntimeError(f"Failed to parse DOCX: {e}")


# ─────────────────────────────────────────────
#  JSON Extraction Helper
# ─────────────────────────────────────────────
def _parse_json_response(text: str) -> dict:
    """Robustly extract JSON from Gemini's response (handles markdown code blocks)."""
    if not text:
        raise ValueError("Empty AI response")
    text = re.sub(r"```(?:json)?\s*", "", text)
    text = re.sub(r"```\s*$", "", text, flags=re.MULTILINE)
    text = text.strip()
    start = text.find("{")
    end = text.rfind("}") + 1
    if start == -1 or end == 0:
        raise ValueError("No JSON object found in AI response")
    return json.loads(text[start:end])


# ─────────────────────────────────────────────
#  Gemini API Helper
# ─────────────────────────────────────────────
def _call_gemini(prompt: str, max_retries: int = 3) -> str:
    """Call Gemini with retry logic and model fallback."""
    if not _configure_gemini():
        raise RuntimeError(
            "Gemini API key not configured. Run: python setup_env.py YOUR_API_KEY"
        )

    last_error = ""
    models = _model_candidates()
    for model_idx, model in enumerate(models):
        for attempt in range(max_retries):
            try:
                response = _gemini_client.models.generate_content(
                    model=model,
                    contents=prompt,
                    config=genai_types.GenerateContentConfig(
                        temperature=0.7,
                        max_output_tokens=8192,
                    ),
                )
                text = response.text
                if text:
                    return text
                last_error = "Empty response from Gemini"
            except Exception as e:
                last_error = str(e)
                if _is_transient_or_model_error(last_error):
                    logger.warning(
                        "Model %s unavailable (%s), trying next fallback", model, last_error
                    )
                    # Brief delay before trying the next model
                    if model_idx < len(models) - 1:
                        time.sleep(1)
                    break  # go to next model
                if _is_permanent_error(last_error):
                    raise RuntimeError(f"Gemini API error: {last_error}")
            if attempt < max_retries - 1:
                time.sleep(2 ** attempt)
    raise RuntimeError(f"Gemini API error after retries: {last_error}")


# ─────────────────────────────────────────────
#  AI Engine Class
# ─────────────────────────────────────────────
class AIEngine:
    def is_ready(self) -> bool:
        return _configure_gemini()

    def using_offline_mode(self) -> bool:
        return not self.is_ready()

    # ── Question Generation ──────────────────
    def generate_questions(self, resume_text: str, job_description: str) -> dict:
        """Generate 7 tailored interview questions from resume + job description."""
        if not _configure_gemini():
            return self._generate_questions_offline(resume_text, job_description)

        resume_snippet = resume_text[:3000]
        jd_snippet = job_description[:1500]

        prompt = f"""You are a senior technical recruiter and interview expert.

RESUME TEXT:
{resume_snippet}

JOB DESCRIPTION:
{jd_snippet}

Analyze the candidate's background and the role requirements. Generate exactly 7 targeted interview questions that will help assess this specific candidate for this specific role.

Question distribution:
- 2 Technical questions (probe skills listed on resume vs role requirements)
- 2 Behavioral questions (STAR-format, based on their actual experience)
- 1 Situational question (role-specific hypothetical scenario)
- 1 Problem-solving question (relevant to the domain)
- 1 Cultural/motivational question (why this role/company)

Return ONLY a valid JSON object with EXACTLY this structure (no markdown, no extra text):
{{
  "candidate_name": "detected name or 'Candidate'",
  "candidate_summary": "2-sentence summary of candidate's background and fit",
  "role": "exact job title from JD",
  "company": "company name from JD or 'the company'",
  "total_questions": 7,
  "questions": [
    {{
      "id": 1,
      "category": "Technical",
      "difficulty": "medium",
      "question": "full question text here",
      "what_to_look_for": "brief note on ideal answer elements",
      "suggested_time_seconds": 120,
      "follow_up": "optional follow-up question"
    }}
  ]
}}

Categories must be one of: Technical, Behavioral, Situational, Problem-Solving, Motivational
Difficulty must be one of: easy, medium, hard"""

        try:
            raw = _call_gemini(prompt)
            result = _parse_json_response(raw)
            if not isinstance(result.get("questions"), list) or not result["questions"]:
                raise ValueError("AI returned invalid question structure")
        except (RuntimeError, ValueError) as exc:
            logger.warning("Question generation failed, using offline mode: %s", exc)
            return self._generate_questions_offline(resume_text, job_description)

        for i, q in enumerate(result["questions"]):
            q.setdefault("id", i + 1)
        result["total_questions"] = len(result["questions"])
        return result

    # ── Answer Analysis ──────────────────────
    def analyze_answer(self, question: str, answer: str,
                       resume_context: str = "", job_context: str = "",
                       biometrics: dict | None = None) -> dict:
        """Score and analyze a candidate's answer."""
        if not answer or not answer.strip():
            return self._empty_analysis("No answer provided")

        if not _configure_gemini():
            result = self._analyze_answer_offline(question, answer, job_context)
            return self._merge_biometrics(result, biometrics)

        prompt = f"""You are an expert interview coach evaluating a candidate's answer.

INTERVIEW QUESTION:
{question}

CANDIDATE'S ANSWER:
{answer}

CONTEXT:
Resume background: {resume_context[:500] if resume_context else 'Not provided'}
Job role: {job_context[:300] if job_context else 'Not provided'}

Evaluate the answer across 5 dimensions and provide detailed coaching feedback.

Return ONLY a valid JSON object with EXACTLY this structure:
{{
  "score": 78,
  "grade": "Good",
  "breakdown": {{
    "relevance": 80,
    "depth": 75,
    "structure": 85,
    "specificity": 70,
    "communication": 82
  }},
  "summary": "One-sentence overall assessment of the answer",
  "strengths": [
    "Specific strength 1",
    "Specific strength 2"
  ],
  "improvements": [
    "Specific area to improve 1",
    "Specific area to improve 2"
  ],
  "model_answer_hint": "2-3 sentence description of what an ideal answer would include",
  "keywords_found": ["keyword1", "keyword2"],
  "keywords_missing": ["important_keyword1"],
  "word_count": 0,
  "answer_quality": "too_short|adequate|good|detailed"
}}

Scoring guide:
- 90-100: Exceptional (comprehensive, structured, highly relevant, specific examples)
- 80-89: Excellent (good depth, relevant, some specifics)
- 70-79: Good (mostly relevant, some depth missing)
- 60-69: Fair (relevant but lacks depth/specifics)
- Below 60: Needs Work (off-topic, too vague, or very short)

Grade must match score: Exceptional/Excellent/Good/Fair/Needs Work"""

        try:
            raw = _call_gemini(prompt)
            result = _parse_json_response(raw)
            result["score"] = max(0, min(100, int(result.get("score", 0))))
        except (RuntimeError, ValueError, TypeError) as exc:
            logger.warning("Answer analysis failed, using offline mode: %s", exc)
            result = self._analyze_answer_offline(question, answer, job_context)
            return self._merge_biometrics(result, biometrics)

        result["word_count"] = len(answer.split())
        result["grade"] = self._score_to_grade(result["score"])
        result.setdefault("breakdown", {})
        result.setdefault("strengths", [])
        result.setdefault("improvements", [])
        return self._merge_biometrics(result, biometrics)

    def _merge_biometrics(self, result: dict, biometrics: dict | None) -> dict:
        if not biometrics:
            return result
        facial = biometrics.get("facial") if isinstance(biometrics.get("facial"), dict) else None
        voice = biometrics.get("voice") if isinstance(biometrics.get("voice"), dict) else None
        if not facial and not voice:
            return result

        fused = fuse_biometrics(facial, voice)
        result["content_score"] = result.get("score", 0)
        result["score"] = adjust_answer_score(result["content_score"], fused)
        result["biometrics"] = fused
        result["grade"] = self._score_to_grade(result["score"])

        delivery_note = (
            f"Presentation delivery scored {fused['presentation_score']}/100 "
            f"({fused['dominant_state']}). "
        )
        if fused.get("recommendations"):
            result.setdefault("improvements", [])
            for tip in fused["recommendations"][:2]:
                if tip not in result["improvements"]:
                    result["improvements"].append(tip)
        result["summary"] = delivery_note + result.get("summary", "")

        if isinstance(result.get("breakdown"), dict):
            result["breakdown"]["communication"] = min(
                100,
                round(result["breakdown"].get("communication", 70) * 0.6
                      + fused["presentation_score"] * 0.4),
            )
        return result

    @staticmethod
    def _score_to_grade(score: int) -> str:
        if score >= 90:
            return "Exceptional"
        if score >= 80:
            return "Excellent"
        if score >= 70:
            return "Good"
        if score >= 60:
            return "Fair"
        return "Needs Work"

    # ── Final Session Report ─────────────────
    def generate_final_report(self, session_data: dict) -> dict:
        """Generate a comprehensive coaching report from the full session."""
        qa_summary = []
        for i, qa in enumerate(session_data.get("answers", []), 1):
            if qa.get("skipped"):
                continue
            analysis = qa.get("analysis") or {}
            qa_summary.append({
                "q_num": i,
                "question_id": qa.get("question_id", qa.get("id", i)),
                "category": qa.get("category", "General"),
                "question": (qa.get("question") or "")[:200],
                "score": analysis.get("score", 0),
                "grade": analysis.get("grade", "N/A"),
            })

        overall = (
            sum(q["score"] for q in qa_summary) // len(qa_summary)
            if qa_summary else 0
        )

        if not _configure_gemini():
            return self._generate_final_report_offline(session_data, overall, qa_summary)

        prompt = f"""You are an expert interview coach. A candidate just completed a mock interview.

SESSION SUMMARY:
- Role applied for: {session_data.get('role', 'Not specified')}
- Total questions answered: {len(qa_summary)}
- Overall average score: {overall}/100
- Duration: {session_data.get('duration', 'N/A')}

QUESTION-BY-QUESTION SCORES:
{json.dumps(qa_summary, indent=2)}

DETAILED ANSWER DATA:
{json.dumps(session_data.get('answers', [])[:5], indent=2, default=str)[:2000]}

Generate a comprehensive coaching report. Return ONLY a valid JSON object:
{{
  "overall_score": {overall},
  "grade": "Good",
  "hiring_probability": "Medium",
  "executive_summary": "3-4 sentence honest assessment of the candidate's interview performance",
  "top_strengths": [
    "Specific strength 1 with evidence",
    "Specific strength 2 with evidence",
    "Specific strength 3 with evidence"
  ],
  "critical_improvements": [
    "Most important area to improve with specific advice",
    "Second area to improve with specific advice",
    "Third area"
  ],
  "action_plan": [
    {{"action": "Specific actionable task", "timeline": "This week", "priority": "high"}},
    {{"action": "Practice task", "timeline": "Next 2 weeks", "priority": "medium"}},
    {{"action": "Long-term improvement", "timeline": "Next month", "priority": "low"}}
  ],
  "best_question_id": 1,
  "weakest_question_id": 2,
  "consistency_score": 75,
  "readiness_level": "Ready|Almost Ready|Needs More Prep",
  "interview_tips": [
    "Tip 1 specific to this candidate",
    "Tip 2",
    "Tip 3"
  ]
}}

Hiring probability must be: High (score>80) / Medium (score 65-80) / Low (score<65)
Readiness level must be: Ready (>80) / Almost Ready (65-80) / Needs More Prep (<65)"""

        try:
            raw = _call_gemini(prompt)
            result = _parse_json_response(raw)
        except (RuntimeError, ValueError) as exc:
            logger.warning("Final report failed, using offline mode: %s", exc)
            return self._generate_final_report_offline(session_data, overall, qa_summary)

        result["overall_score"] = overall
        result["grade"] = self._score_to_grade(overall)
        if qa_summary:
            result["best_question_id"] = max(qa_summary, key=lambda q: q["score"])["question_id"]
            result["weakest_question_id"] = min(qa_summary, key=lambda q: q["score"])["question_id"]
        return result

    # ── Offline fallbacks (no API key / API failure) ───────
    def _extract_role(self, job_description: str) -> str:
        for pattern in (
            r"(?:role|position|title)\s*[:\-]\s*([^\n]+)",
            r"(?:hiring|seeking|looking for)\s+(?:a\s+)?([^\n.]{5,60})",
            r"^([A-Z][A-Za-z\s/&-]{4,50}(?:Engineer|Developer|Manager|Analyst|Designer|Scientist))",
        ):
            match = re.search(pattern, job_description, re.I | re.M)
            if match:
                return match.group(1).strip()[:60]
        first_line = next((ln.strip() for ln in job_description.splitlines() if ln.strip()), "")
        return first_line[:60] or "Target Role"

    def _extract_name(self, resume_text: str) -> str:
        lines = [ln.strip() for ln in resume_text.splitlines() if ln.strip()]
        if not lines:
            return "Candidate"
        first = lines[0]
        if len(first.split()) <= 4 and not any(c in first for c in "@|:"):
            return first
        return "Candidate"

    def _extract_skills(self, text: str) -> list[str]:
        catalog = [
            "python", "java", "javascript", "react", "node", "sql", "aws", "azure",
            "docker", "kubernetes", "machine learning", "data analysis", "communication",
            "leadership", "agile", "rest", "api", "flask", "django", "tensorflow",
        ]
        lower = text.lower()
        return [s.title() for s in catalog if s in lower][:6]

    def _generate_questions_offline(self, resume_text: str, job_description: str) -> dict:
        role = self._extract_role(job_description)
        name = self._extract_name(resume_text)
        skills = self._extract_skills(resume_text + " " + job_description)
        skill_str = ", ".join(skills) if skills else "relevant technical skills"

        templates = [
            ("Technical", "medium", f"Walk me through a project where you used {skill_str}. What was your role and outcome?"),
            ("Technical", "hard", f"How would you approach a challenging problem in a {role} role using {skill_str}?"),
            ("Behavioral", "medium", "Tell me about a time you worked under pressure and delivered results. Use the STAR format."),
            ("Behavioral", "medium", "Describe a conflict with a teammate and how you resolved it professionally."),
            ("Situational", "medium", f"You join as {role} and inherit a delayed project. What do you do in the first two weeks?"),
            ("Problem-Solving", "hard", "How would you debug a production issue when logs are incomplete and stakeholders need a quick update?"),
            ("Motivational", "easy", f"Why are you interested in this {role} position, and what makes you a strong fit?"),
        ]
        questions = [
            {
                "id": i,
                "category": category,
                "difficulty": difficulty,
                "question": question,
                "what_to_look_for": "Specific examples, clear structure, measurable outcomes",
                "suggested_time_seconds": 120,
                "follow_up": "Can you quantify the impact?",
            }
            for i, (category, difficulty, question) in enumerate(templates, 1)
        ]
        return {
            "candidate_name": name,
            "candidate_summary": (
                f"{name} appears to match a {role} profile with experience in {skill_str}. "
                "Questions were generated in demo mode - add a valid Gemini API key for deeper personalization."
            ),
            "role": role,
            "company": "the company",
            "total_questions": len(questions),
            "questions": questions,
            "offline_mode": True,
        }

    def _analyze_answer_offline(self, question: str, answer: str, job_context: str = "") -> dict:
        words = answer.split()
        word_count = len(words)
        lower = answer.lower()

        score = 35
        if word_count >= 30:
            score += 15
        if word_count >= 80:
            score += 15
        if word_count >= 120:
            score += 10
        if any(k in lower for k in ("because", "result", "impact", "learned", "achieved")):
            score += 10
        if any(k in lower for k in ("situation", "task", "action", "star")):
            score += 8
        if re.search(r"\d+%|\d+\s*(users|days|months|years|percent)", lower):
            score += 7
        score = min(92, score)
        grade = self._score_to_grade(score)

        def clamp(v: int) -> int:
            return max(0, min(100, v))

        skills = self._extract_skills(job_context + " " + question)
        found = [s for s in skills if s.lower() in lower]
        missing = [s for s in skills if s.lower() not in lower][:3]

        return {
            "score": score,
            "grade": grade,
            "breakdown": {
                "relevance": clamp(score + 5),
                "depth": clamp(score - 5 + (10 if word_count > 60 else 0)),
                "structure": clamp(score + (8 if "first" in lower or "then" in lower else 0)),
                "specificity": clamp(score + (10 if found else -5)),
                "communication": clamp(score),
            },
            "summary": (
                "Demo analysis (AI unavailable). "
                + ("Solid length and structure." if word_count >= 60 else "Answer is short - add more detail and examples.")
            ),
            "strengths": [
                "Clear written response" if word_count >= 40 else "Attempted the question",
                "Uses relevant context" if found else "Room to tie answer to role requirements",
            ],
            "improvements": [
                "Add a valid Gemini API key in Settings for AI-powered feedback",
                "Use STAR format with measurable results",
                "Mention specific tools, teams, and outcomes",
            ],
            "model_answer_hint": "Describe context, your actions, and quantifiable results in 100-200 words.",
            "keywords_found": found,
            "keywords_missing": missing,
            "word_count": word_count,
            "answer_quality": "too_short" if word_count < 30 else "adequate" if word_count < 80 else "good",
            "offline_mode": True,
        }

    def _generate_final_report_offline(self, session_data: dict, overall: int,
                                       qa_summary: list | None = None) -> dict:
        qa_summary = qa_summary or []
        hire = "High" if overall >= 80 else "Medium" if overall >= 65 else "Low"
        readiness = "Ready" if overall >= 80 else "Almost Ready" if overall >= 65 else "Needs More Prep"
        best_id = max(qa_summary, key=lambda q: q["score"])["question_id"] if qa_summary else 1
        weak_id = min(qa_summary, key=lambda q: q["score"])["question_id"] if qa_summary else 1

        return {
            "overall_score": overall,
            "grade": self._score_to_grade(overall),
            "hiring_probability": hire,
            "executive_summary": (
                f"You completed {len(session_data.get('answers', []))} questions with an average score of {overall}/100. "
                "This report used demo mode. Add a valid Gemini API key in Settings for full AI coaching."
            ),
            "top_strengths": [
                "Completed the mock interview session",
                "Provided answers across multiple question types",
                "Engaged with structured feedback",
            ],
            "critical_improvements": [
                "Configure Gemini API key for personalized AI analysis",
                "Add more quantified outcomes in answers",
                "Practice STAR-format behavioral responses",
            ],
            "action_plan": [
                {"action": "Add Gemini API key via Settings banner", "timeline": "Today", "priority": "high"},
                {"action": "Re-run interview with AI analysis enabled", "timeline": "This week", "priority": "medium"},
                {"action": "Practice 2 questions daily with specific examples", "timeline": "Next 2 weeks", "priority": "low"},
            ],
            "best_question_id": best_id,
            "weakest_question_id": weak_id,
            "consistency_score": overall,
            "readiness_level": readiness,
            "interview_tips": [
                "Use the STAR method for behavioral questions",
                "Include numbers and outcomes where possible",
                "Connect answers back to the job description",
            ],
            "offline_mode": True,
        }

    # ── Helpers ──────────────────────────────
    def _empty_analysis(self, reason: str) -> dict:
        return {
            "score": 0,
            "grade": "Needs Work",
            "breakdown": {"relevance": 0, "depth": 0, "structure": 0, "specificity": 0, "communication": 0},
            "summary": reason,
            "strengths": [],
            "improvements": ["Please provide a complete answer"],
            "model_answer_hint": "",
            "keywords_found": [],
            "keywords_missing": [],
            "word_count": 0,
            "answer_quality": "too_short",
        }